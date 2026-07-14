"""
FORGE File Generation Integration
===================================
Generate PDF, DOCX, XLSX, CSV, and Markdown files from agent output.

Uses only stdlib + optional lightweight libraries:
  PDF   — reportlab (pip install reportlab) — falls back to HTML if absent
  DOCX  — python-docx (pip install python-docx)
  XLSX  — openpyxl (pip install openpyxl)
  CSV   — stdlib csv module (always available)
  MD    — plain text write (always available)

All outputs go to the configured workspace directory.

Env vars
--------
  FORGE_WORKSPACE   output directory (default: workspace)
"""
from __future__ import annotations

import csv
import io
import json
import os
from pathlib import Path
from typing import Any, Dict, List, Optional, Union

from forge.telemetry.logger import get_logger

log = get_logger("files")

def _workspace() -> Path:
    """Return current workspace root from live config."""
    try:
        from forge.config import get_config
        return Path(get_config().workspace_dir)
    except Exception:
        return Path(os.environ.get("FORGE_WORKSPACE", "workspace")).resolve()


# ── Result type ───────────────────────────────────────────────────────────────

class FileResult:
    def __init__(self, success: bool, path: Optional[Path] = None, error: str = ""):
        self.success = success
        self.path = path
        self.error = error

    def __str__(self):
        if self.success:
            return f"Written: {self.path}"
        return f"Error: {self.error}"


# ── Markdown / plain text ─────────────────────────────────────────────────────

def write_markdown(content: str, filename: str, subdir: str = "") -> FileResult:
    """Write markdown content to workspace."""
    out = _workspace() / subdir / filename if subdir else _workspace() / filename
    out.parent.mkdir(parents=True, exist_ok=True)
    try:
        out.write_text(content, encoding="utf-8")
        log.info("Markdown written", path=str(out))
        return FileResult(success=True, path=out)
    except Exception as e:
        return FileResult(success=False, error=str(e))


# ── CSV ───────────────────────────────────────────────────────────────────────

def write_csv(
    rows: List[Dict[str, Any]],
    filename: str,
    subdir: str = "",
    fieldnames: Optional[List[str]] = None,
) -> FileResult:
    """Write a list of dicts to CSV."""
    out = _workspace() / subdir / filename if subdir else _workspace() / filename
    out.parent.mkdir(parents=True, exist_ok=True)
    try:
        fields = fieldnames or (list(rows[0].keys()) if rows else [])
        with open(out, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
            writer.writeheader()
            writer.writerows(rows)
        log.info("CSV written", path=str(out), rows=len(rows))
        return FileResult(success=True, path=out)
    except Exception as e:
        return FileResult(success=False, error=str(e))


# ── XLSX ──────────────────────────────────────────────────────────────────────

def write_xlsx(
    sheets: Dict[str, List[Dict[str, Any]]],
    filename: str,
    subdir: str = "",
) -> FileResult:
    """
    Write an Excel workbook. sheets is {sheet_name: [row_dict, ...]}.
    Requires openpyxl.
    """
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        return FileResult(success=False,
                          error="openpyxl not installed. Run: pip install openpyxl")

    out = _workspace() / subdir / filename if subdir else _workspace() / filename
    out.parent.mkdir(parents=True, exist_ok=True)

    try:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # remove default sheet

        header_fill = PatternFill("solid", fgColor="2B579A")
        header_font = Font(color="FFFFFF", bold=True)

        for sheet_name, rows in sheets.items():
            ws = wb.create_sheet(title=sheet_name[:31])  # Excel limit
            if not rows:
                continue

            headers = list(rows[0].keys())
            # Write header row
            for col_idx, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_idx, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center")

            # Write data rows
            for row_idx, row in enumerate(rows, 2):
                for col_idx, header in enumerate(headers, 1):
                    ws.cell(row=row_idx, column=col_idx, value=row.get(header, ""))

            # Auto-size columns
            for col_idx, _ in enumerate(headers, 1):
                col_letter = get_column_letter(col_idx)
                max_len = max(
                    len(str(ws.cell(row=r, column=col_idx).value or ""))
                    for r in range(1, len(rows) + 2)
                )
                ws.column_dimensions[col_letter].width = min(max_len + 4, 50)

        wb.save(out)
        log.info("XLSX written", path=str(out), sheets=list(sheets.keys()))
        return FileResult(success=True, path=out)
    except Exception as e:
        return FileResult(success=False, error=str(e))


# ── DOCX ──────────────────────────────────────────────────────────────────────

def write_docx(
    content: Union[str, List[Dict[str, Any]]],
    filename: str,
    title: str = "",
    subdir: str = "",
) -> FileResult:
    """
    Write a Word document.
    content can be:
      - str: treated as markdown-like text (# heading, ## subheading, plain text)
      - list of dicts: [{type: heading|subheading|paragraph|bullet|code, text: ...}]
    Requires python-docx.
    """
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        return FileResult(success=False,
                          error="python-docx not installed. Run: pip install python-docx")

    out = _workspace() / subdir / filename if subdir else _workspace() / filename
    out.parent.mkdir(parents=True, exist_ok=True)

    try:
        doc = Document()

        # Document title
        if title:
            doc.add_heading(title, level=0)

        # Parse content
        if isinstance(content, str):
            blocks = _parse_markdown_blocks(content)
        else:
            blocks = content

        for block in blocks:
            btype = block.get("type", "paragraph")
            text  = block.get("text", "")

            if btype == "heading":
                doc.add_heading(text, level=1)
            elif btype == "subheading":
                doc.add_heading(text, level=2)
            elif btype == "bullet":
                doc.add_paragraph(text, style="List Bullet")
            elif btype == "numbered":
                doc.add_paragraph(text, style="List Number")
            elif btype == "code":
                p = doc.add_paragraph()
                run = p.add_run(text)
                run.font.name = "Courier New"
                run.font.size = Pt(9)
            elif btype == "table" and isinstance(text, list) and text:
                headers = list(text[0].keys())
                table = doc.add_table(rows=1, cols=len(headers))
                table.style = "Table Grid"
                hdr_cells = table.rows[0].cells
                for i, h in enumerate(headers):
                    hdr_cells[i].text = h
                for row_data in text[1:]:
                    row_cells = table.add_row().cells
                    for i, h in enumerate(headers):
                        row_cells[i].text = str(row_data.get(h, ""))
            else:
                doc.add_paragraph(text)

        doc.save(out)
        log.info("DOCX written", path=str(out))
        return FileResult(success=True, path=out)
    except Exception as e:
        return FileResult(success=False, error=str(e))


def _parse_markdown_blocks(text: str) -> List[Dict[str, Any]]:
    """Convert markdown-like text into block dicts for write_docx."""
    blocks = []
    in_code = False
    code_lines = []

    for line in text.splitlines():
        if line.startswith("```"):
            if in_code:
                blocks.append({"type": "code", "text": "\n".join(code_lines)})
                code_lines = []
                in_code = False
            else:
                in_code = True
            continue

        if in_code:
            code_lines.append(line)
            continue

        if line.startswith("# "):
            blocks.append({"type": "heading", "text": line[2:].strip()})
        elif line.startswith("## "):
            blocks.append({"type": "subheading", "text": line[3:].strip()})
        elif line.startswith("### "):
            blocks.append({"type": "subheading", "text": line[4:].strip()})
        elif line.startswith("- ") or line.startswith("* "):
            blocks.append({"type": "bullet", "text": line[2:].strip()})
        elif line.strip():
            blocks.append({"type": "paragraph", "text": line.strip()})

    return blocks


# ── PDF ───────────────────────────────────────────────────────────────────────

def write_pdf(
    content: str,
    filename: str,
    title: str = "",
    subdir: str = "",
) -> FileResult:
    """
    Write a PDF. Uses reportlab if available, otherwise writes an HTML
    file that can be printed to PDF from a browser.
    """
    out = _workspace() / subdir / filename if subdir else _workspace() / filename
    out.parent.mkdir(parents=True, exist_ok=True)

    # Try reportlab first
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Preformatted
        from reportlab.lib.enums import TA_LEFT

        doc = SimpleDocTemplate(str(out), pagesize=A4,
                                leftMargin=20*mm, rightMargin=20*mm,
                                topMargin=20*mm, bottomMargin=20*mm)
        styles = getSampleStyleSheet()
        story = []

        if title:
            story.append(Paragraph(title, styles["Title"]))
            story.append(Spacer(1, 6*mm))

        for line in content.splitlines():
            if line.startswith("# "):
                story.append(Paragraph(line[2:], styles["Heading1"]))
            elif line.startswith("## "):
                story.append(Paragraph(line[3:], styles["Heading2"]))
            elif line.strip():
                story.append(Paragraph(line, styles["Normal"]))
            else:
                story.append(Spacer(1, 3*mm))

        doc.build(story)
        log.info("PDF written via reportlab", path=str(out))
        return FileResult(success=True, path=out)

    except ImportError:
        pass

    # Fallback: write HTML with print stylesheet
    html_out = out.with_suffix(".html")
    escaped = content.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    html = f"""<!DOCTYPE html>
<html><head><meta charset="utf-8">
<title>{title}</title>
<style>
  body {{ font-family: Georgia, serif; max-width: 800px; margin: 40px auto; line-height: 1.6; }}
  h1 {{ color: #1a1a2e; }} h2 {{ color: #16213e; }}
  pre {{ background: #f5f5f5; padding: 12px; border-radius: 4px; overflow-x: auto; }}
  @media print {{ body {{ margin: 0; max-width: 100%; }} }}
</style></head>
<body>
{'<h1>' + title + '</h1>' if title else ''}
<pre style="white-space: pre-wrap;">{escaped}</pre>
</body></html>"""
    try:
        html_out.write_text(html, encoding="utf-8")
        log.info("PDF fallback: HTML written", path=str(html_out))
        return FileResult(success=True, path=html_out,
                          error="reportlab not installed — wrote HTML instead. "
                                "Open in browser and Print → Save as PDF. "
                                "Install reportlab for native PDF: pip install reportlab")
    except Exception as e:
        return FileResult(success=False, error=str(e))